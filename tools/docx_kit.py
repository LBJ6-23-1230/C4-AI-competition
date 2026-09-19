"""知学 Mate · Word 文档生成共享样式库。

从 `generate_task_plan_docx.py` 抽取，供本目录下 5 份下一步工作文档复用，
保证所有交付文档的字体、配色、标题层级、表格样式完全一致。

依赖：python-docx
用法：
    from docx_kit import (Document, add_title, h, para, rich, bullets,
                          table, note_box, page_break, COLOR_*)
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# --------------------------------------------------------------------------- 常量
FONT_CN = "微软雅黑"
FONT_EN = "Segoe UI"

COLOR_PRIMARY = RGBColor(0x1F, 0x4E, 0x79)   # 主色（深蓝，与工程品牌 #2563EB 同系）
COLOR_RED = RGBColor(0xC0, 0x00, 0x00)       # 阻塞 / 红线
COLOR_ORANGE = RGBColor(0xB4, 0x53, 0x09)    # 风险 / 未完成
COLOR_GREEN = RGBColor(0x1E, 0x6B, 0x3A)     # 已完成 / 正向
COLOR_GREY = RGBColor(0x59, 0x59, 0x59)      # 辅助说明
COLOR_BLUE_LIGHT = RGBColor(0x2E, 0x5C, 0x8A)

# 交付文档统一输出到工程 deliverables/（换机器也能用）
OUT_DIR = Path(__file__).resolve().parent.parent / "deliverables"


# --------------------------------------------------------------------------- 基础样式
def _set_run_font(run, size=10.5, bold=False, color=None, cn=FONT_CN, en=FONT_EN):
    run.font.name = en
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    specs = {
        "Heading 1": (16, COLOR_PRIMARY),
        "Heading 2": (13.5, COLOR_PRIMARY),
        "Heading 3": (11.5, COLOR_BLUE_LIGHT),
        "Heading 4": (10.5, RGBColor(0x40, 0x40, 0x40)),
    }
    for name, (size, color) in specs.items():
        style = doc.styles[name]
        style.font.name = FONT_EN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)

    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)


def new_document() -> Document:
    doc = Document()
    style_document(doc)
    return doc


# --------------------------------------------------------------------------- 构件
def add_title(doc: Document, text: str, subtitle: str | None = None) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(p.add_run(text), size=20, bold=True, color=COLOR_PRIMARY)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(p2.add_run(subtitle), size=11, color=COLOR_GREY)


def h(doc: Document, text: str, level: int = 1):
    return doc.add_heading(text, level=level)


def para(doc: Document, text: str, size: float = 10.5, bold: bool = False,
         color=None, indent: bool = False, align=None):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.6)
    if align is not None:
        p.alignment = align
    _set_run_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def rich(doc: Document, segments, indent: bool = False, size: float = 10.5):
    """segments = [(text, bold, color), ...]"""
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.6)
    for text, bold, color in segments:
        _set_run_font(p.add_run(text), size=size, bold=bold, color=color)
    return p


def bullets(doc: Document, items, style: str = "List Bullet") -> None:
    for item in items:
        if isinstance(item, tuple):
            text, bold = item
        else:
            text, bold = item, False
        p = doc.add_paragraph(style=style)
        _set_run_font(p.add_run(text), size=10.5, bold=bold)


def code(doc: Document, text: str) -> None:
    """等宽代码块（用 Consolas，浅灰底纹）。"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    _set_run_font(p.add_run(text), size=9.5, cn="Consolas", en="Consolas")
    pPr = p._p.get_or_add_pPr()
    shd = pPr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F4F7")
    pPr.append(shd)


def table(doc: Document, headers, rows, widths=None, font_size: float = 9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_run_font(p.add_run(text), size=font_size, bold=True,
                      color=RGBColor(0xFF, 0xFF, 0xFF))
        shd = hdr[i]._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {})
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), "1F4E79")
        hdr[i]._tc.get_or_add_tcPr().append(shd)
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            _set_run_font(p.add_run(str(text)), size=font_size)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph()
    return t


def note_box(doc: Document, text: str, color=COLOR_RED, label: str = "关键提醒"):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    _set_run_font(p.add_run(f"【{label}】"), size=10.5, bold=True, color=color)
    _set_run_font(p.add_run(text), size=10.5, color=color)
    return p


def page_break(doc: Document) -> None:
    doc.add_page_break()


def save(doc: Document, filename: str) -> Path:
    """保存文档，并**显式检测文件被占用**的情况。

    为什么需要：交付文档经常被 Word / WPS 打开着，此时 python-docx 的 save()
    会抛 PermissionError。若调用方把输出重定向丢弃，就会**静默保留旧文件**，
    表现为"明明改了数据，文档却没更新"——本项目实际踩过一次。

    这里改为：先探测能否独占写入；被占用时抛出带明确指引的异常，
    而不是留下一个内容过期的文件。
    """
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / filename

    if path.exists():
        try:
            probe = open(path, "r+b")
            probe.close()
        except PermissionError:
            raise PermissionError(
                f"无法写入 {path}\n"
                f"  原因：该文件正被 Word / WPS 等程序打开（写锁）。\n"
                f"  处理：请先关闭它，然后重新运行本生成脚本。\n"
                f"  注意：继续运行不会更新该文件，只会保留旧内容。"
            ) from None

    doc.save(str(path))
    return path
